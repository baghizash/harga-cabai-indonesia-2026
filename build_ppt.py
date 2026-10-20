#!/usr/bin/env python3
"""Build Portofolio-Analisis-Harga-Cabai.pptx (16:9, 10 slides, ID).

Redesign design-taste: Tema T1 (marine / bright red / butter), varian
Editorial Bold. Konten dan semua angka identik dengan deck lama;
sumber angka: findings.json (data PIHPS / Bank Indonesia).
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
F = json.loads((HERE / "findings.json").read_text())
CHARTS = HERE / "charts"
OUTS = [
    HERE / "Portofolio-Analisis-Harga-Cabai.pptx",
    Path.home() / "workspace" / "your_files" / "Portofolio-Analisis-Harga-Cabai.pptx",
]

# Tema T1
GROUND = RGBColor(0xF7, 0xF4, 0xEE)
INK = RGBColor(0x0B, 0x2A, 0x44)
MARINE = RGBColor(0x0E, 0x4D, 0x7A)
RED = RGBColor(0xE5, 0x32, 0x2D)
BUTTER = RGBColor(0xF5, 0xC5, 0x42)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Liberation Serif"
SANS = "Noto Sans"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rp(v):
    return f"Rp{v:,}".replace(",", ".")


def new_slide(bg=GROUND):
    s = prs.slides.add_slide(blank)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return s


def textbox(slide, l, t, w, h):
    tx = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False,
            font=SANS, align=None, space_after=8, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    if line is not None:
        p.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16, color=INK, first=False, space_after=12):
    p = add_par(tf, text, size, color, first=first, space_after=space_after, line=1.12)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    return p


def rect(slide, l, t, w, h, color, line_color=None):
    sh = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def title_block(slide, text):
    tf = textbox(slide, 0.7, 0.32, 11.93, 0.85)
    add_par(tf, text, 31, INK, bold=True, first=True, font=SERIF)
    rect(slide, 0.72, 1.16, 2.3, 0.055, RED)


def bullets_box(slide, bullets, l, t, w, h, size=16):
    tf = textbox(slide, l, t, w, h)
    for i, b in enumerate(bullets):
        add_bullet(tf, b, size=size, first=(i == 0))
    return tf


vol = F["volatilitas_2026"]
rt = F["rata_tahunan"]
lonj = F["lonjakan_oktober_rawit_merah"]
prov = F["provinsi_2026"]
most_vol = F["komoditas_paling_volatile"]
rm = prov["Cabai Rawit Merah"]
cov = F["coverage"]

# ---- Slide 1: Cover (Data-Hero) ----
s = new_slide()
rect(s, 0, 0, 0.16, 7.5, RED)
tf = textbox(s, 0.85, 0.95, 5.15, 1.85)
add_par(tf, "ANALISIS HARGA CABAI INDONESIA 2026", 37, INK, bold=True,
        first=True, font=SERIF, line=1.05)
tf = textbox(s, 0.85, 2.95, 5.15, 0.85)
add_par(tf, "Lonjakan Oktober, Volatilitas, dan Disparitas Antar Provinsi",
        17, MARINE, bold=True, first=True, line=1.15)
rect(s, 0.87, 3.92, 1.6, 0.05, RED)
tf = textbox(s, 0.85, 4.2, 5.15, 1.3)
add_par(tf, "Data harian PIHPS / Bank Indonesia - 4 komoditas cabai, 34 provinsi, "
            "Januari 2025 sampai Oktober 2026", 13.5, INK, first=True, line=1.2)
add_par(tf, "Plus: paket dashboard Power BI (star schema, DAX, panduan rakit)",
        13.5, INK, line=1.2)
rect(s, 0.87, 6.28, 5.0, 0.03, MARINE)
tf = textbox(s, 0.85, 6.48, 5.15, 0.5)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 16, INK, bold=True, first=True)
# Hero: chart lonjakan Oktober + panel angka jangkar
s.shapes.add_picture(str(CHARTS / "05_lonjakan_oktober.png"),
                     Inches(6.35), Inches(0.7), width=Inches(6.35))
rect(s, 6.35, 3.95, 6.35, 2.62, MARINE)
tf = textbox(s, 6.7, 4.22, 5.7, 0.4)
add_par(tf, "PUNCAK CABAI RAWIT MERAH - 12 OKTOBER 2026", 12, BUTTER,
        bold=True, first=True)
tf = textbox(s, 6.7, 4.58, 5.7, 1.25)
p = tf.paragraphs[0]
r = p.add_run(); r.text = rp(F["puncak_2026"]["Cabai Rawit Merah"]["harga"])
r.font.name = SANS; r.font.size = Pt(63); r.font.bold = True; r.font.color.rgb = BUTTER
r = p.add_run(); r.text = "/kg"
r.font.name = SANS; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = BUTTER
tf = textbox(s, 6.7, 5.95, 5.7, 0.4)
add_par(tf, "Harga nasional tertinggi Cabai Rawit Merah sepanjang 2026",
        12.5, WHITE, first=True)

# ---- Slide 2: Latar belakang ----
s = new_slide()
title_block(s, "Kenapa Cabai?")
bullets_box(s, [
    "Cabai adalah komoditas pangan paling volatile di Indonesia: harganya bisa "
    "melonjak puluhan persen dalam hitungan hari dan langsung terasa di inflasi.",
    "Analisis ini memakai 4 komoditas PIHPS: Cabai Merah Besar, Cabai Merah Keriting, "
    "Cabai Rawit Hijau, dan Cabai Rawit Merah.",
    "Pertanyaan: bagaimana harga bergerak di 2026, cabai mana yang paling liar, "
    "di provinsi mana harganya paling mahal, dan seberapa besar lonjakan Oktober 2026?",
    "Jawaban dicari dari data harian resmi, bukan dari perkiraan.",
], 0.7, 1.55, 11.93, 4.1, size=16.5)
rect(s, 0.7, 5.7, 11.93, 1.0, BUTTER)
for i, (num, lab) in enumerate([("4", "komoditas cabai"), ("34", "provinsi"),
                                ("Jan 2025 - Okt 2026", "periode data harian")]):
    tf = textbox(s, 1.05 + i * 4.0, 5.86, 3.8, 0.75)
    add_par(tf, num, 17, INK, bold=True, first=True, space_after=1)
    add_par(tf, lab, 11.5, INK, space_after=0)
tf = textbox(s, 0.7, 6.88, 11.93, 0.4)
add_par(tf, "Sumber data: API publik PIHPS / Bank Indonesia (hargapangan.id), Pasar Tradisional.",
        11.5, MARINE, italic=True, first=True)

# ---- Slide 3: Data & metode ----
s = new_slide()
title_block(s, "Data dan Metode")
bullets_box(s, [
    f"Data harian API PIHPS: {cov['tanggal_awal']} sampai {cov['tanggal_akhir']}, "
    f"{cov['baris_nasional']:,} baris nasional + {cov['baris_provinsi_total']:,} baris provinsi "
    f"({cov['baris_provinsi']:,} dengan harga).",
    "Setiap pertanyaan dijawab dengan query SQL (sql/queries.sql) dan divisualisasikan "
    "dalam 10 chart dari data yang sama.",
    "Untuk Power BI, data disusun ulang menjadi star schema: 1 tabel fakta + "
    "3 tabel dimensi, lengkap dengan measure DAX dan panduan merakit dashboard.",
    "Catatan jujur: PIHPS tidak mempublikasikan seri provinsi untuk beberapa kombinasi "
    f"provinsi x komoditas ({len(F['gaps'])} kombinasi) - dicatat sebagai gap, tidak diisi angka lain.",
], 0.7, 1.55, 7.2, 5.2, size=16)
rect(s, 8.25, 1.5, 4.38, 4.95, MARINE)
tf = textbox(s, 8.6, 1.85, 3.75, 4.4)
stats = [
    (f"{cov['baris_nasional']:,}", "baris nasional"),
    (f"{cov['baris_provinsi_total']:,}", "baris provinsi"),
    (f"{cov['baris_provinsi']:,}", "baris provinsi dengan harga"),
    (f"{len(F['gaps'])}", "kombinasi provinsi x komoditas tanpa seri (gap)"),
]
for i, (num, lab) in enumerate(stats):
    add_par(tf, num, 27, BUTTER, bold=True, first=(i == 0), space_after=1)
    add_par(tf, lab, 12.5, WHITE, space_after=15)

# ---- Slide 4: Tren nasional ----
s = new_slide()
title_block(s, "Tren Nasional 2025-2026")
bullets_box(s, [
    f"Rata-rata nasional 2026 tertinggi: Cabai Rawit Merah {rp(rt['Cabai Rawit Merah']['2026'])}/kg, "
    f"naik {rt['Cabai Rawit Merah']['perubahan_persen']}% dibanding 2025.",
    f"Cabai Merah Besar rata-rata {rp(rt['Cabai Merah Besar']['2026'])}/kg "
    f"({rt['Cabai Merah Besar']['perubahan_persen']:+}% vs 2025); Keriting "
    f"{rp(rt['Cabai Merah Keriting']['2026'])}/kg ({rt['Cabai Merah Keriting']['perubahan_persen']:+}%); "
    f"Rawit Hijau {rp(rt['Cabai Rawit Hijau']['2026'])}/kg ({rt['Cabai Rawit Hijau']['perubahan_persen']:+}%).",
    "Keempat komoditas bergerak searah pada lonjakan besar, tapi dengan amplitudo berbeda.",
], 0.7, 1.55, 5.3, 5.2, size=16)
s.shapes.add_picture(str(CHARTS / "01_tren_nasional_4cabai.png"),
                     Inches(6.25), Inches(1.6), width=Inches(6.55))
rect(s, 6.25, 4.9, 6.55, 1.05, BUTTER)
tf = textbox(s, 6.55, 5.12, 6.0, 0.65)
add_par(tf, f"Cabai Rawit Merah rata-rata {rp(rt['Cabai Rawit Merah']['2026'])}/kg di 2026, "
            f"naik {rt['Cabai Rawit Merah']['perubahan_persen']}% dibanding 2025.",
        14.5, INK, bold=True, first=True, line=1.12)

# ---- Slide 5: Lonjakan Oktober ----
s = new_slide()
title_block(s, "Lonjakan Oktober 2026")
bullets_box(s, [
    (f"Cabai Rawit Merah menembus Rp100.000/kg pada {F['tembus_rawit_merah']['100000']['tanggal']} "
     f"({rp(F['tembus_rawit_merah']['100000']['harga'])}/kg).") if "100000" in F["tembus_rawit_merah"]
    else "Ambang Rp100.000/kg tidak tertembus pada seri nasional dalam periode data.",
    (f"Dan menembus Rp105.000/kg pada {F['tembus_rawit_merah']['105000']['tanggal']} "
     f"({rp(F['tembus_rawit_merah']['105000']['harga'])}/kg).") if "105000" in F["tembus_rawit_merah"]
    else "Ambang Rp105.000/kg tidak tertembus pada seri nasional dalam periode data.",
    f"Dari akhir September ({rp(lonj['akhir_september']['harga'])}/kg) ke puncak Oktober "
    f"({rp(lonj['maks_oktober']['harga'])}/kg pada {lonj['maks_oktober']['tanggal']}), "
    f"naik {lonj['kenaikan_persen']}%.",
    f"Kenaikan terbesar terjadi pada September: rata-rata bulanan naik "
    f"{F['mom_rawit_merah_2026']['2026-09']:+}% dibanding Agustus, lalu Oktober naik lagi "
    f"{F['mom_rawit_merah_2026']['2026-10']:+}%.",
    "Lonjakan ini yang membuat cabai menjadi berita utama pertengahan Oktober 2026.",
], 0.7, 1.5, 5.5, 5.4, size=15.5)
s.shapes.add_picture(str(CHARTS / "05_lonjakan_oktober.png"),
                     Inches(6.5), Inches(1.55), width=Inches(6.3))
rect(s, 6.5, 4.68, 6.3, 1.92, MARINE)
tf = textbox(s, 6.82, 4.9, 5.7, 0.95)
p = tf.paragraphs[0]
r = p.add_run(); r.text = rp(lonj["maks_oktober"]["harga"])
r.font.name = SANS; r.font.size = Pt(38); r.font.bold = True; r.font.color.rgb = BUTTER
r = p.add_run(); r.text = "/kg"
r.font.name = SANS; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BUTTER
tf = textbox(s, 6.82, 5.85, 5.7, 0.6)
add_par(tf, f"Puncak {lonj['maks_oktober']['tanggal']} - naik {lonj['kenaikan_persen']}% "
            f"dari akhir September ({rp(lonj['akhir_september']['harga'])}/kg).",
        12.5, WHITE, first=True, line=1.15)

# ---- Slide 6: Volatilitas ----
s = new_slide()
title_block(s, "Volatilitas: Cabai Mana yang Paling Liar?")
s.shapes.add_picture(str(CHARTS / "03_volatilitas.png"),
                     Inches(0.7), Inches(1.6), width=Inches(5.7))
rect(s, 0.7, 4.72, 5.7, 1.0, RED)
tf = textbox(s, 1.0, 4.95, 5.15, 0.6)
add_par(tf, f"Paling volatile: {most_vol} - CV {vol[most_vol]['cv_persen']}%",
        15, WHITE, bold=True, first=True, line=1.12)
bullets_box(s, [
    f"{most_vol} paling volatile di 2026: koefisien variasi {vol[most_vol]['cv_persen']}%, "
    f"rentang {rp(vol[most_vol]['min'])}/kg sampai {rp(vol[most_vol]['maks'])}/kg.",
    f"Yang paling stabil: {F['komoditas_paling_stabil']} (CV {vol[F['komoditas_paling_stabil']]['cv_persen']}%).",
    "Volatilitas tinggi berarti harga sulit diprediksi: pedagang dan konsumen "
    "sama-sama menanggung ketidakpastian.",
    "Lonjakan harian terbesar 2026: "
    + ", ".join(f"{j['komoditas']} {j['perubahan_persen']:+}% pada {j['tanggal']}"
                for j in F["lonjakan_harian_terbesar_2026"][:2]) + ".",
], 6.75, 1.55, 5.9, 5.2, size=16)

# ---- Slide 7: Disparitas provinsi ----
s = new_slide()
title_block(s, "Disparitas Antar Provinsi")
bullets_box(s, [
    f"2026, Cabai Rawit Merah termahal di {rm['termahal']['provinsi']} "
    f"(rata-rata {rp(rm['termahal']['harga'])}/kg), termurah di {rm['termurah']['provinsi']} "
    f"({rp(rm['termurah']['harga'])}/kg) - selisih {rp(rm['selisih'])}/kg "
    f"({rm['selisih_persen']}%).",
    f"DKI Jakarta {rp(F['jakarta_vs_nasional_2026']['Cabai Rawit Merah']['jakarta'])}/kg, "
    f"{F['jakarta_vs_nasional_2026']['Cabai Rawit Merah']['selisih_persen']:+}% "
    "terhadap rata-rata nasional.",
    "Harga di provinsi timur dan kepulauan cenderung lebih mahal: biaya logistik "
    "dan rantai pasok yang lebih panjang.",
], 0.7, 1.55, 5.4, 3.5, size=16)
rect(s, 0.7, 5.2, 5.4, 1.0, BUTTER)
tf = textbox(s, 1.0, 5.42, 4.85, 0.6)
add_par(tf, f"Selisih termahal-termurah Rawit Merah: {rm['selisih_persen']}%",
        15, INK, bold=True, first=True, line=1.12)
s.shapes.add_picture(str(CHARTS / "06_provinsi_termahal.png"),
                     Inches(6.35), Inches(1.55), width=Inches(6.45))

# ---- Slide 8: Power BI (bento asimetris) ----
s = new_slide()
title_block(s, "Paket Dashboard Power BI")
panels = [
    (0.7, 1.5, 7.3, 2.45, MARINE, WHITE,
     "Star schema siap import: fact_harga_cabai + dim_komoditas + dim_wilayah + dim_tanggal."),
    (8.25, 1.5, 4.38, 2.45, BUTTER, INK,
     "Layout 4 halaman: Overview, Per Komoditas, Per Wilayah, Volatilitas - lengkap "
     "dengan spesifikasi tiap visual."),
    (0.7, 4.15, 4.38, 2.6, RED, WHITE,
     "Panduan rakit langkah demi langkah untuk Power BI Desktop di Windows: import, "
     "relasi, mark date table, tempel DAX, bangun visual, publish."),
    (5.33, 4.15, 7.3, 2.6, GROUND, INK,
     "15 measure DAX siap tempel: rata-rata, MoM %, YoY %, MA 7 hari, volatilitas (CV %), "
     "peringkat provinsi, selisih vs nasional."),
]
for (l, t, w, h, bg, fg, text) in panels:
    if bg == GROUND:
        rect(s, l, t, w, h, GROUND, line_color=MARINE)
    else:
        rect(s, l, t, w, h, bg)
    tf = textbox(s, l + 0.3, t + 0.28, w - 0.6, h - 0.5)
    add_par(tf, text, 15.5, fg, first=True, line=1.18)

# ---- Slide 9: Insight & rekomendasi ----
top_jump = F["lonjakan_harian_terbesar_2026"][0]
s = new_slide()
title_block(s, "Insight dan Rekomendasi")
rect(s, 0.7, 1.55, 0.09, 5.15, RED)
items = [
    f"Pantau Rawit Merah (paling mahal, rata-rata 2026 {rp(rt['Cabai Rawit Merah']['2026'])}/kg) "
    f"dan Keriting (paling volatile, CV {vol['Cabai Merah Keriting']['cv_persen']}%).",
    f"Siapkan skenario lonjakan: lonjakan harian terbesar {top_jump['perubahan_persen']:+}% "
    f"({top_jump['komoditas']}, {top_jump['tanggal']}) menunjukkan harga bisa melonjak dalam sehari.",
    "Fokus ke provinsi timur: disparitas harga antar provinsi mencapai puluhan persen - "
    "distribusi dan logistik adalah bagian dari masalah.",
    "Gunakan dashboard Power BI untuk pemantauan harian: MA 7 hari dan MoM % memberi "
    "sinyal lebih awal daripada harga harian mentah.",
]
tf = textbox(s, 1.1, 1.6, 11.55, 5.2)
for i, text in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(20)
    p.line_spacing = 1.15
    r = p.add_run(); r.text = f"{i + 1}.  "
    r.font.name = SANS; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = RED
    r = p.add_run(); r.text = text
    r.font.name = SANS; r.font.size = Pt(16.5); r.font.color.rgb = INK

# ---- Slide 10: Penutup (Designed Close) ----
s = new_slide(bg=MARINE)
rect(s, 0.9, 0.85, 1.8, 0.07, BUTTER)
tf = textbox(s, 0.9, 1.1, 11.5, 1.0)
add_par(tf, "Terima kasih.", 40, WHITE, bold=True, first=True, font=SERIF)
tf = textbox(s, 0.9, 2.25, 11.3, 0.9)
add_par(tf, "Deck ini terbuka untuk didiskusikan. Semua angka di deck ini dihitung dari "
            "data harian API PIHPS / Bank Indonesia - dapat direproduksi dengan menjalankan "
            "scripts/fetch_pihps.py lalu scripts/analyze.py.",
        16, GROUND, first=True, line=1.25)
tf = textbox(s, 0.9, 3.55, 11.5, 0.85)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 33, BUTTER, bold=True, first=True, font=SERIF)
tf = textbox(s, 0.9, 4.5, 11.5, 0.5)
add_par(tf, "github.com/baghizash/harga-cabai-indonesia-2026", 16, WHITE,
        bold=True, first=True)
tf = textbox(s, 0.9, 5.75, 11.5, 1.3)
add_par(tf, "Struktur proyek: analysis.ipynb, data/raw, sql/queries.sql, scripts, charts, "
            "powerbi/ (kit dashboard), findings.json.", 11.5, GROUND, first=True, line=1.2)
add_par(tf, "Keterbatasan: PIHPS melaporkan pada hari kerja pasar (bukan setiap hari kalender); "
            "beberapa kombinasi provinsi x komoditas tidak dipublikasikan dan dicatat sebagai gap.",
        11.5, GROUND, line=1.2)

for out in OUTS:
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"PPT saved: {out}")
